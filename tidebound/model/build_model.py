"""Builds model/revenue-model.xlsx from shared/config/catalog.json.
Inputs (blue) are copied from config; every output is a live Excel formula.
Run: python3 model/build_model.py  (then optionally: python3 <xlsx skill>/scripts/recalc.py model/revenue-model.xlsx)"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
cat = json.load(open(os.path.join(ROOT, "shared", "config", "catalog.json")))

ARIAL = "Arial"
BLUE = Font(name=ARIAL, color="0000FF")
BLACK = Font(name=ARIAL, color="000000")
GREEN = Font(name=ARIAL, color="008000")
BOLD = Font(name=ARIAL, bold=True)
TITLE = Font(name=ARIAL, bold=True, size=14)
KEY = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="DDEBF7")
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
USD4 = '$#,##0.0000;($#,##0.0000);"-"'
PCT = '0.0%;(0.0%);"-"'
INT = '#,##0;(#,##0);"-"'
thin = Side(style="thin", color="BBBBBB")

wb = Workbook()
inp = wb.active
inp.title = "Inputs"
inp.column_dimensions["A"].width = 44
for col, w in zip("BCDEFG", [16, 14, 14, 14, 16, 16]):
    inp.column_dimensions[col].width = w

def label(ws, cell, text, font=BLACK):
    ws[cell] = text
    ws[cell].font = font

def value(ws, cell, v, fmt=None, font=BLUE, note=None, key=False):
    ws[cell] = v
    ws[cell].font = font
    if fmt:
        ws[cell].number_format = fmt
    if note:
        ws[cell].comment = Comment(note, "Tidebound")
    if key:
        ws[cell].fill = KEY

label(inp, "A1", "Tidebound revenue model - inputs", TITLE)
label(inp, "A2", "Blue cells are inputs (copied from shared/config/catalog.json; edit here to explore). Yellow = key assumptions. Black cells are formulas.")
inp["A2"].alignment = Alignment(wrap_text=True)

label(inp, "A4", "Daily players (DAU)", BOLD)
value(inp, "B4", cat["planning"]["dailyPlayers"], INT, key=True, note="Planning assumption from the brief (catalog.json planning.dailyPlayers).")

# Product table
label(inp, "A6", "Products", BOLD)
headers = ["Product", "Price (USD)", "Monthly buy rate", "Robux price", "Spend per DAU ($)", "Store type"]
for i, h in enumerate(headers):
    c = inp.cell(row=7, column=1 + i, value=h)
    c.font = BOLD
    c.fill = HEAD
first = 8
products = cat["products"]
for i, p in enumerate(products):
    r = first + i
    label(inp, f"A{r}", p["id"])
    value(inp, f"B{r}", cat["priceTiers"][p["priceTier"]]["usd"], USD2)
    value(inp, f"C{r}", cat["planning"]["buyRates"].get(p["id"], 0), PCT, key=True,
          note="Share of daily players who buy this per month. Planning assumption from the brief; gem bundles have no planning rate (0)." if i == 0 else None)
    inp[f"E{r}"] = f"=B{r}*C{r}"
    inp[f"E{r}"].number_format = USD4
    label(inp, f"F{r}", p["robloxType"])
last = first + len(products) - 1

r0 = last + 2  # fees block
label(inp, f"A{r0}", "Standalone app platform fees", BOLD)
fees = cat["platformFees"]
rows = [
    ("iOS commission (Small Business Program)", fees["ios"]["commissionRate"], PCT, "Apple Small Business Program rate (brief). Enrolment is on the store checklist."),
    ("Google Play rate, first $1M per year", fees["googlePlay"]["rateUnderThreshold"], PCT, "Brief: 15% on the first $1M of annual earnings."),
    ("Google Play rate above $1M per year", fees["googlePlay"]["rateOverThreshold"], PCT, "Google's standard rate above the threshold."),
    ("Google Play threshold ($ per year)", fees["googlePlay"]["thresholdUsdPerYear"], USD, None),
    ("Platform mix: iOS share of spend", fees["platformMix"]["ios"], PCT, "Assumption: 50/50. Both stores charge 15% at launch scale, so the mix does not change the default result."),
    ("Platform mix: Google share of spend", fees["platformMix"]["googlePlay"], PCT, None),
]
fee_cells = {}
for i, (t, v, f, n) in enumerate(rows):
    r = r0 + 1 + i
    label(inp, f"A{r}", t)
    value(inp, f"B{r}", v, f, note=n)
    fee_cells[i] = f"Inputs!$B${r}"

rr = r0 + 8  # Roblox block
for i in range(len(products)):
    r = first + i
    inp[f"D{r}"] = f"=ROUND(B{r}/$B${rr + 1},0)"
    inp[f"D{r}"].number_format = INT
label(inp, f"A{rr}", "Roblox payouts", BOLD)
rb = cat["roblox"]
rows2 = [
    ("USD per Robux (price mapping)", rb["usdPerRobux"], '0.000', "Brief: Robux price = USD price / 0.01."),
    ("Roblox platform fee", rb["platformFee"], PCT, "Brief: Roblox keeps 30% of each sale."),
    ("DevEx cash-out ($ per earned Robux)", rb["devExUsdPerRobux"], '0.0000', "Brief: $0.0038 per earned Robux."),
    ("Premium payouts ($ per DAU per month)", cat["planning"]["premiumPayoutsPerDailyPlayerMonthlyUsd"], USD4, "Engagement-based Premium Payouts. Unknown before launch; 0 keeps the planning comparison."),
]
rob = {}
for i, (t, v, f, n) in enumerate(rows2):
    r = rr + 1 + i
    label(inp, f"A{r}", t)
    value(inp, f"B{r}", v, f, note=n)
    rob[i] = f"Inputs!$B${r}"

rc = rr + 6
label(inp, f"A{rc}", "Fixed monthly costs", BOLD)
rows3 = [
    ("Standalone: backend", cat["standalone"]["backendCostMonthlyUsd"], "Brief default: $400/month."),
    ("Standalone: marketing", cat["planning"]["marketingMonthlyUsd"], "Brief default: $1,000/month."),
    ("Roblox: companion backend", rb["companionBackendMonthlyUsd"], "Default 0 to match the brief's comparison. A small host for /backend costs about $5-20/month."),
    ("Roblox: marketing (sponsored ads)", rb["marketingMonthlyUsd"], "Default 0 to match the brief's comparison, which gives Roblox no fixed costs."),
]
cost = {}
for i, (t, v, n) in enumerate(rows3):
    r = rc + 1 + i
    label(inp, f"A{r}", t)
    value(inp, f"B{r}", v, USD, note=n)
    cost[i] = f"Inputs!$B${r}"

# ---------------- Model sheet
m = wb.create_sheet("Model")
m.column_dimensions["A"].width = 46
m.column_dimensions["B"].width = 18
m.column_dimensions["C"].width = 18
m.column_dimensions["D"].width = 40
label(m, "A1", "Tidebound revenue model - outputs (all formulas)", TITLE)
for i, h in enumerate(["Metric", "Standalone app", "Roblox", "Formula"]):
    c = m.cell(row=3, column=1 + i, value=h)
    c.font = BOLD
    c.fill = HEAD

dau = "Inputs!$B$4"
spend = f"SUM(Inputs!$E${first}:$E${last})"
lines = [
    # (label, standalone formula, roblox formula, fmt, explanation)
    ("Daily players", f"={dau}", f"={dau}", INT, "Input"),
    ("Spend per daily player per month", f"={spend}", f"={spend}", USD4, "Sum of price x monthly buy rate"),
    ("Gross monthly player spend", "=B5*B6", "=C5*C6", USD, "DAU x spend per DAU"),
    ("Gross annual player spend", "=B7*12", "=C7*12", USD, "Monthly x 12"),
    ("Google Play rate applied", f"=IF(B8*{fee_cells[5]}>{fee_cells[3]},{fee_cells[2]},{fee_cells[1]})", None, PCT, "15% until Google spend passes $1M/year"),
    ("Platform share of spend", f"={fee_cells[0]}*{fee_cells[4]}+B9*{fee_cells[5]}", f"=1-C11", PCT, "Standalone: blended store fee. Roblox: 1 - keep rate"),
    ("Developer keep rate", "=1-B10", f"=(1-{rob[1]})*{rob[2]}/{rob[0]}", PCT, "Roblox: (1 - 30%) x DevEx $ per Robux / $ per Robux"),
    ("Platform fee ($ per month)", "=B7*B10", "=C7*C10", USD, "Gross x platform share"),
    ("Premium Payouts ($ per month)", "=0", f"={rob[3]}*C5", USD, "Roblox only"),
    ("Net revenue per month", "=B7-B12+B13", "=C7-C12+C13", USD, "Gross - fee + payouts"),
    ("Net revenue per year", "=B14*12", "=C14*12", USD, "Monthly x 12"),
    ("Fixed costs per month", f"={cost[0]}+{cost[1]}", f"={cost[2]}+{cost[3]}", USD, "Backend + marketing"),
    ("Profit per month", "=B14-B16", "=C14-C16", USD, "Net - fixed costs"),
    ("Profit per year", "=B17*12", "=C17*12", USD, "Monthly x 12"),
    ("Net margin per daily player per month", "=B6*B11", f"=C6*C11+{rob[3]}", USD4, "Spend per DAU x keep rate (+ payouts)"),
    ("Break-even daily players", "=IF(B19>0,B16/B19,0)", "=IF(C19>0,C16/C19,0)", INT, "Fixed costs / margin per DAU"),
]
for i, (t, fs, fr, fmt, expl) in enumerate(lines):
    r = 5 + i - 1 + 1  # rows 5..20
    r = 4 + i + 1
    label(m, f"A{r}", t)
    for col, f in (("B", fs), ("C", fr)):
        if f is not None:
            m[f"{col}{r}"] = f
            m[f"{col}{r}"].number_format = fmt
            m[f"{col}{r}"].font = GREEN if f.startswith("=Inputs!") else BLACK
    label(m, f"D{r}", expl)
assert m["A20"].value == "Break-even daily players"

label(m, "A22", "Crossover: daily players where standalone profit overtakes Roblox", BOLD)
m["B22"] = "=IF(B19>C19,(B16-C16)/(B19-C19),0)"
m["B22"].number_format = INT
m["B22"].fill = KEY
label(m, "D22", "(Standalone fixed - Roblox fixed) / (standalone margin - Roblox margin)")

# Planning check
label(m, "A24", "Planning check (numbers from the brief)", BOLD)
for i, h in enumerate(["Brief", "Model", "Difference"]):
    c = m.cell(row=25, column=2 + i, value=h)
    c.font = BOLD
    c.fill = HEAD
checks = [
    ("Spend per daily player per month", 0.15, "=B6", USD4),
    ("Gross per month at 10,000 DAU", 1506, "=B7", USD),
    ("Net of 15% store cut per month", 1280, "=B14", USD),
    ("Fixed costs per month", 1400, "=B16", USD),
    ("Standalone break-even DAU", 11000, "=B20", INT),
    ("Roblox keep rate", 0.27, "=C11", PCT),
    ("Crossover DAU", 16000, "=B22", INT),
]
for i, (t, brief, f, fmt) in enumerate(checks):
    r = 26 + i
    label(m, f"A{r}", t)
    value(m, f"B{r}", brief, fmt, note="Value stated in the brief.")
    m[f"C{r}"] = f
    m[f"C{r}"].number_format = fmt
    m[f"D{r}"] = f"=C{r}-B{r}"
    m[f"D{r}"].number_format = fmt
label(m, "A34", "Differences: break-even is 10,936 (brief rounds to ~11,000); Roblox keep is 26.6% (0.7 x 0.0038 / 0.01; brief rounds to 27%); crossover is ~15,900 (brief ~16,000) and assumes Roblox has no fixed costs.")
m["A34"].alignment = Alignment(wrap_text=True)
m.row_dimensions[34].height = 45

# Scaling table
t = wb.create_sheet("Scenarios")
t.column_dimensions["A"].width = 16
for col in "BCDEF":
    t.column_dimensions[col].width = 20
label(t, "A1", "Profit by daily players (uses Model and Inputs)", TITLE)
for i, h in enumerate(["Daily players", "Gross / month", "Standalone profit / month", "Roblox profit / month", "Better option"]):
    c = t.cell(row=3, column=1 + i, value=h)
    c.font = BOLD
    c.fill = HEAD
for i, d in enumerate([1000, 2500, 5000, 7500, 10000, 12500, 15000, 16000, 20000, 25000, 30000, 40000, 50000, 75000, 100000]):
    r = 4 + i
    value(t, f"A{r}", d, INT)
    t[f"B{r}"] = f"=A{r}*Model!$B$6"
    t[f"C{r}"] = f"=B{r}*Model!$B$11-Model!$B$16"
    t[f"D{r}"] = f"=B{r}*Model!$C$11+A{r}*Inputs!$B${rr+4}-Model!$C$16"
    t[f"E{r}"] = f'=IF(C{r}>D{r},"Standalone","Roblox")'
    for col in "BCD":
        t[f"{col}{r}"].number_format = USD
        t[f"{col}{r}"].font = GREEN

notes = wb.create_sheet("Notes")
notes.column_dimensions["A"].width = 110
lines = [
    "How to use: change blue cells on Inputs. Every output on Model and Scenarios recalculates.",
    "Source of defaults: shared/config/catalog.json (prices, buy rates, fees, Roblox mapping, costs). Regenerate with python3 model/build_model.py after editing config.",
    "Gem bundles have no planning buy rate in the brief, so they contribute 0 by default. Add rates on Inputs to include them.",
    "Roblox keep rate = (1 - platform fee) x DevEx rate / USD per Robux = 0.7 x 0.0038 / 0.01 = 26.6% of player spend.",
    "Roblox Subscriptions pay out differently from Robux sales; the model applies the same keep rate to Harbor Club. Check the current Roblox subscription payout terms before relying on it.",
    "The brief's crossover (~16,000 DAU) assumes Roblox has no fixed costs. With the same $1,000 marketing on both, the crossover falls to about 4,500 DAU (see Inputs Roblox costs).",
    "Tidebound ships on Roblox. With $1,000/month marketing on Roblox, break-even there is about 25,000 DAU at the planning buy rates.",
]
for i, l in enumerate(lines):
    notes[f"A{i+1}"] = l
    notes[f"A{i+1}"].font = BLACK
    notes[f"A{i+1}"].alignment = Alignment(wrap_text=True)

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != ARIAL:
                c.font = Font(name=ARIAL, bold=c.font.bold, size=c.font.size, color=c.font.color)
wb.save(os.path.join(ROOT, "model", "revenue-model.xlsx"))
print("wrote model/revenue-model.xlsx")

# ---------------- HTML calculator with the same defaults
defaults = {
    "dailyPlayers": cat["planning"]["dailyPlayers"],
    "products": [{"id": p["id"], "usd": cat["priceTiers"][p["priceTier"]]["usd"], "rate": cat["planning"]["buyRates"].get(p["id"], 0)} for p in products],
    "ios": fees["ios"]["commissionRate"], "googleUnder": fees["googlePlay"]["rateUnderThreshold"], "googleOver": fees["googlePlay"]["rateOverThreshold"],
    "googleThreshold": fees["googlePlay"]["thresholdUsdPerYear"], "mixIos": fees["platformMix"]["ios"], "mixGoogle": fees["platformMix"]["googlePlay"],
    "usdPerRobux": rb["usdPerRobux"], "robloxFee": rb["platformFee"], "devEx": rb["devExUsdPerRobux"], "premium": cat["planning"]["premiumPayoutsPerDailyPlayerMonthlyUsd"],
    "standaloneBackend": cat["standalone"]["backendCostMonthlyUsd"], "standaloneMarketing": cat["planning"]["marketingMonthlyUsd"],
    "robloxBackend": rb["companionBackendMonthlyUsd"], "robloxMarketing": rb["marketingMonthlyUsd"],
}
template = open(os.path.join(ROOT, "model", "calculator_template.html")).read()
open(os.path.join(ROOT, "model", "revenue-model.html"), "w").write(template.replace("/*DEFAULTS*/null", json.dumps(defaults)))
print("wrote model/revenue-model.html")
